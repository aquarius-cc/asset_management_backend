"""操作日志导出行级安全测试（P0，不可回退）。

【锁定的不变量】
导出的可见范围必须**恒等于**列表接口的可见范围，即 Selector 层
``_scope_by_user`` 的三态语义：

- ``get_department_codes_for_user`` 返回 ``None``（超级管理员 / 审计员 /
  无 Employee 记录）-> 不限部门
- 返回**空列表**（部门级角色但未挂部门）-> 零行
- 返回**部门列表** -> 仅本部门（含子部门）资产的操作记录

任一环节漏掉 ``_scope_by_user``，即为跨部门数据泄露（与已修 P1-3 同类）。
因此本文件以"解析导出 xlsx 内容"的方式断言，而非仅断言状态码。
"""

import io

import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.assetmanagement.models import Asset, AssetType, Storage
from apps.assetmanagement.models.operation_log import AssetOperationLog
from apps.assetmanagement.operation_log_views import AssetOperationLogExportView
from apps.usermanagement.models import Department, Employee, EmployeeRole


User = get_user_model()

TEST_PASSWORD = "Test@12345"
EXPORT_URL = reverse("operation-log-export")
_phone_seq = iter(range(1, 10_000))

#: 需 transaction=True：``StreamingExcelResponse.close()`` 会触发
#: ``HttpResponseBase.close()`` 发送 ``request_finished``，其接收方
#: ``close_old_connections`` 会关闭数据库连接，与 pytest-django 默认的
#: 原子事务包裹冲突（表现为 "the connection is closed"）。
pytestmark = pytest.mark.django_db(transaction=True)


def _phone():
    return f"138{next(_phone_seq):08d}"


def _make_user(username, role, department=None):
    """创建带角色的 AuthUser + Employee（Employee 决定行级部门范围）。"""
    user = User.objects.create_user(
        auth_username=username,
        password=TEST_PASSWORD,
        auth_is_staff=(role == EmployeeRole.SYSTEM_ADMIN),
    )
    Employee.objects.create(
        employee_jobcode=username,
        employee_name=f"{username}员工",
        employee_department=department,
        role=role,
        employee_phone=_phone(),
    )
    return user


def _make_storage(storage_code, department):
    """仓库自身无部门字段，部门归属经 storage_manager(Employee) 关联。"""
    manager = Employee.objects.create(
        employee_jobcode=f"WH-{storage_code}",
        employee_name=f"{department.department_name}仓管",
        employee_department=department,
        role=EmployeeRole.ASSET_ADMIN,
        employee_phone=_phone(),
    )
    return Storage.objects.create(
        storage_code=storage_code,
        storage_name=f"{department.department_name}仓库",
        storage_address=f"{department.department_name}地址",
        storage_manager=manager,
    )


def _make_asset(asset_code, department, storage, type_code):
    asset_type = AssetType.objects.create(type_code=type_code, type_name=f"类型{type_code}")
    keeper = Employee.objects.create(
        employee_jobcode=f"KP-{type_code}",
        employee_name=f"{department.department_name}保管人",
        employee_department=department,
        role=EmployeeRole.ASSET_ADMIN,
        employee_phone=_phone(),
    )
    return Asset.objects.create(
        asset_code=asset_code,
        asset_name=f"{department.department_name}资产",
        asset_purchase_price=1000,
        asset_purchase_date="2024-01-01",
        asset_entry_date="2024-01-15",
        asset_type_recordcode=asset_type,
        asset_storage_recordcode=storage,
        asset_entry_person_recordcode=keeper,
        asset_current_status="in_store",
    )


def _make_log(asset):
    return AssetOperationLog.objects.create(
        asset_code=asset.asset_code,
        asset_name=asset.asset_name,
        operation_type=AssetOperationLog.OperationType.UPDATE,
        description=f"{asset.asset_code} 更新操作",
    )


@pytest.fixture
def dept_a():
    return Department.objects.create(department_code="DEPT-A", department_name="A部门")


@pytest.fixture
def dept_b():
    return Department.objects.create(department_code="DEPT-B", department_name="B部门")


@pytest.fixture
def two_dept_assets(dept_a, dept_b):
    """A/B 两部门各一台资产，且各有一条操作记录。"""
    asset_a = _make_asset("AST-A001", dept_a, _make_storage("SA-EXP", dept_a), "TA")
    asset_b = _make_asset("AST-B001", dept_b, _make_storage("SB-EXP", dept_b), "TB")
    return _make_log(asset_a), _make_log(asset_b)


def _export_asset_codes(user, query=""):
    """调用导出端点并返回 xlsx 第一列的资产编码集合（排除表头）。"""
    request = APIRequestFactory().get(f"{EXPORT_URL}{query}")
    force_authenticate(request, user=user)
    response = AssetOperationLogExportView.as_view()(request)
    assert response.status_code == 200, response.status_code
    try:
        payload = b"".join(response.streaming_content)
    finally:
        response.close()
    sheet = load_workbook(io.BytesIO(payload))["操作记录"]
    return {
        row[0]
        for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True)
        if row[0] is not None
    }


# --------------------------------------------------------------------------
# 负例：他部门数据不得出现在导出文件中
# --------------------------------------------------------------------------


def test_export_excludes_other_department_rows(two_dept_assets, dept_a):
    """A 部门资产管理员导出时，B 部门的操作记录绝不能出现在文件里。"""
    user = _make_user("exp-a", EmployeeRole.ASSET_ADMIN, dept_a)
    codes = _export_asset_codes(user)
    assert codes == {"AST-A001"}
    assert "AST-B001" not in codes


def test_export_scope_matches_list_scope(two_dept_assets, dept_a):
    """导出行集合必须与列表接口返回的行集合一致（可见性不变量）。"""
    from apps.assetmanagement.selectors.operation_log_selector import OperationLogSelector

    user = _make_user("exp-a2", EmployeeRole.ASSET_ADMIN, dept_a)
    listed = {log.asset_code for log in OperationLogSelector.query_operation_logs(user)}
    assert _export_asset_codes(user) == listed == {"AST-A001"}


def test_export_denied_for_department_role_without_department(two_dept_assets):
    """部门级角色但未挂部门 -> 导出权限最严兜底为 403。

    ``get_user_role`` 对"部门级角色但无部门"返回 ``None``（最严兜底），
    不在 ``_EXPORT_EXCEL_ROLES`` 内，故请求在权限层即被拒。
    """
    request = APIRequestFactory().get(EXPORT_URL)
    user = _make_user("exp-nodept", EmployeeRole.ASSET_ADMIN, None)
    force_authenticate(request, user=user)
    response = AssetOperationLogExportView.as_view()(request)
    assert response.status_code == 403


def test_selector_yields_zero_rows_for_department_role_without_department(two_dept_assets):
    """纵深防御：即便绕过权限层，Selector 的行级过滤也必须返回零行。

    ``get_department_codes_for_user`` 返回空列表 -> ``_scope_by_user``
    走 ``queryset.none()`` 分支，而非放行全部。
    """
    from apps.assetmanagement.selectors.operation_log_selector import OperationLogSelector

    user = _make_user("exp-nodept2", EmployeeRole.ASSET_ADMIN, None)
    queryset = OperationLogSelector.build_operation_logs_queryset(user)
    assert queryset.count() == 0
    assert list(queryset) == []


# --------------------------------------------------------------------------
# 正例：不限部门的角色可见全量
# --------------------------------------------------------------------------


@pytest.mark.parametrize("role", [EmployeeRole.SYSTEM_ADMIN, EmployeeRole.AUDITOR])
def test_export_unrestricted_roles_see_all_departments(two_dept_assets, role):
    user = _make_user(f"exp-{role}", role, None)
    assert _export_asset_codes(user) == {"AST-A001", "AST-B001"}


# --------------------------------------------------------------------------
# 过滤与分批参数在受控可见性内生效
# --------------------------------------------------------------------------


def test_export_respects_filter_within_scope(two_dept_assets, dept_a):
    user = _make_user("exp-a3", EmployeeRole.ASSET_ADMIN, dept_a)
    assert _export_asset_codes(user, "?asset_code=AST-B001") == set()
    assert _export_asset_codes(user, "?asset_code=AST-A001") == {"AST-A001"}


def test_export_limit_offset_still_scoped(two_dept_assets, dept_a):
    """limit/offset 不得成为绕过行级过滤的后门。"""
    user = _make_user("exp-a4", EmployeeRole.ASSET_ADMIN, dept_a)
    assert _export_asset_codes(user, "?limit=1&offset=0") == {"AST-A001"}
    assert _export_asset_codes(user, "?offset=1") == set()
