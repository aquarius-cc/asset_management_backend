"""员工域行级隔离测试（BF-048，P1 安全，不可回退）。

【锁定的不变量】
员工域任何返回**员工行**的端点，可见范围必须**恒等于**调用者的部门范围，即
``core.department_scope.get_department_codes_for_user`` 的三态语义：

- ``None``（superuser / system_admin / auditor / 无 Employee 记录）-> 不限部门
- **空列表**（部门级角色但未挂部门）-> 零行，且写类/导出类权限一并降级
  （``core.permissions.get_user_role`` 降级为 None）
- **部门列表** -> dept_manager 见本部门 + 全部下级部门；asset_admin / regular_user 见本部门

受保护面共 8 个端点：五条主路径（list / retrieve / search / statistics / export）
+ 三个原先**绕过取行口径**的旁路（active_employees / by_auth_user /
get_employee_by_jobcode）。任一环节漏掉行级收窄，即为跨部门员工档案（PII）泄露。

【与 BF-047 的关系】
BF-047 统一了「筛选口径」（本次查询要哪些行），BF-048 补上「权限口径」
（本次调用者能看哪些行）——两者共同构成取行语义，缺一即分叉。故本文件同时
断言 **导出行集合 == 列表行集合**（BF-047 不变量在 RBAC 维度上的延续）。

【本文件不覆盖】
``get_department_by_jobcode`` 返回部门而非员工档案，跨部门部门可见性另立评估。
"""

import io

import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.usermanagement.models import Department, Employee, EmployeeRole, EmployeeStatus


User = get_user_model()

TEST_PASSWORD = "Test@12345"
LIST_URL = reverse("employees-list")
STATISTICS_URL = reverse("employees-statistics")
EXPORT_URL = reverse("employees-export-excel")
ACTIVE_URL = reverse("employees-active-employees")
SEARCH_URL = reverse("employees-global-search")

#: 需 transaction=True：导出端点返回 StreamingExcelResponse，其 close() 会
#: 发送 request_finished，接收方 close_old_connections 会关库，与 pytest-django
#: 默认原子事务包裹冲突（"the connection is closed"）。与
#: test_operation_log_export_scope.py 同因同解。
pytestmark = pytest.mark.django_db(transaction=True)

_phone_seq = iter(range(1, 10_000))


def _phone() -> str:
    return f"138{next(_phone_seq):08d}"


def _make_requester(username: str, role: str, department: Department | None) -> User:
    """创建带角色的 AuthUser + Employee（Employee 决定行级部门范围）。

    ``get_employee_for_user`` 以 ``AuthUser.auth_username == Employee.employee_jobcode``
    隐式关联（非 FK），故两者 jobcode 必须同名。
    """
    user = User.objects.create_user(
        auth_username=username,
        password=TEST_PASSWORD,
        auth_phone=_phone(),
        auth_is_staff=(role == EmployeeRole.SYSTEM_ADMIN),
    )
    Employee.objects.create(
        employee_jobcode=username,
        employee_name=f"{username}本人",
        employee_department=department,
        role=role,
        employee_phone=_phone(),
    )
    return user


def _make_employee(jobcode: str, name: str, department: Department | None) -> Employee:
    """被查看的数据员工（不需要登录态）。"""
    return Employee.objects.create(
        employee_jobcode=jobcode,
        employee_name=name,
        employee_department=department,
        role=EmployeeRole.REGULAR_USER,
        employee_status=EmployeeStatus.ACTIVE,
        employee_phone=_phone(),
    )


def _get(user: User, url: str):
    """以指定用户身份 GET 目标端点（走真实路由，故 action 映射与生产一致）。"""
    client = APIClient()
    client.force_authenticate(user=user)
    return client.get(url)


def _rows(response) -> list[dict]:
    """分页响应 -> 员工行列表（兼容分页 / 不分页两种形态）。"""
    data = response.data["data"]
    return data["results"] if "results" in data else data


def _jobcodes(response) -> set[str]:
    return {row["employee_jobcode"] for row in _rows(response)}


def _export_jobcodes(user: User) -> set[str]:
    """调用导出端点并返回 xlsx 中的工号集合（第一列为员工工号）。"""
    response = _get(user, EXPORT_URL)
    assert response.status_code == 200, response.status_code
    try:
        payload = b"".join(response.streaming_content)
    finally:
        response.close()
    sheet = load_workbook(io.BytesIO(payload)).worksheets[0]
    return {row[0] for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True) if row[0] is not None}


# --------------------------------------------------------------------------
# 夹具：总部(ROOT) -> 下级(CHILD) 层级 + 无关部门(OTHER)，三方各 1 名员工
# --------------------------------------------------------------------------


@pytest.fixture
def dept_root():
    return Department.objects.create(department_code="DEPT-ROOT", department_name="总部", path="/DEPT-ROOT")


@pytest.fixture
def dept_child(dept_root):
    """总部下级部门。``get_all_descendants`` 走 path 前缀匹配，故 path 须形如 /ROOT/CHILD。"""
    return Department.objects.create(
        department_code="DEPT-CHILD",
        department_name="下级部门",
        parent=dept_root,
        path="/DEPT-ROOT/DEPT-CHILD",
        level=1,
    )


@pytest.fixture
def dept_other():
    return Department.objects.create(department_code="DEPT-OTHER", department_name="无关部门", path="/DEPT-OTHER")


@pytest.fixture
def three_employees(dept_root, dept_child, dept_other):
    """三方各 1 名在职员工；返回 (总部, 下级, 无关) 三元组。"""
    return (
        _make_employee("EMP-ROOT", "总部员工", dept_root),
        _make_employee("EMP-CHILD", "下级员工", dept_child),
        _make_employee("EMP-OTHER", "无关员工", dept_other),
    )


# --------------------------------------------------------------------------
# 三态语义
# --------------------------------------------------------------------------


def test_dept_manager_sees_own_department_and_descendants(dept_root, three_employees):
    """dept_manager 见本部门 + 全部下级部门，无关部门不可见。

    期望集合含调用者自身的员工行（``_make_requester`` 会在其部门建 Employee），
    这不是泄露——本人档案本就在可见范围内。
    """
    user = _make_requester("MGR-ROOT", EmployeeRole.DEPT_MANAGER, dept_root)
    response = _get(user, LIST_URL)
    assert response.status_code == 200, response.status_code
    assert _jobcodes(response) == {"MGR-ROOT", "EMP-ROOT", "EMP-CHILD"}


def test_regular_user_sees_only_own_department(dept_child, three_employees):
    """regular_user 见本部门，不含上级部门与下级部门。"""
    user = _make_requester("USR-CHILD", EmployeeRole.REGULAR_USER, dept_child)
    assert _jobcodes(_get(user, LIST_URL)) == {"USR-CHILD", "EMP-CHILD"}


@pytest.mark.parametrize("role", [EmployeeRole.SYSTEM_ADMIN, EmployeeRole.AUDITOR])
def test_global_roles_are_unscoped(dept_other, three_employees, role):
    """system_admin / auditor 为全局角色，不因部门字段缺失或范围而收权。"""
    user = _make_requester(f"U-{role}", role, dept_other)
    assert _jobcodes(_get(user, LIST_URL)) == {f"U-{role}", "EMP-ROOT", "EMP-CHILD", "EMP-OTHER"}


def test_dept_role_without_department_sees_nothing(three_employees):
    """部门级角色但未挂部门 -> 零行（最严兜底），且导出权限一并降级为 403。"""
    user = _make_requester("MGR-NODEPT", EmployeeRole.DEPT_MANAGER, None)
    response = _get(user, LIST_URL)
    assert response.status_code == 200, response.status_code
    assert _rows(response) == []
    # CanExportExcel 经 get_user_role 降级为 None -> 403，与数据范围兜底同源
    assert _get(user, EXPORT_URL).status_code == 403


# --------------------------------------------------------------------------
# statistics：聚合口径必须等于可见行范围（否则数字与列表对不上）
# --------------------------------------------------------------------------


def test_statistics_aggregates_scoped_rows_only(dept_root, three_employees):
    user = _make_requester("MGR-ROOT", EmployeeRole.DEPT_MANAGER, dept_root)
    listed = _jobcodes(_get(user, LIST_URL))
    stats = _get(user, STATISTICS_URL).data["data"]
    assert stats["total_employees"] == len(listed)
    assert stats["total_employees"] == 3  # MGR-ROOT + EMP-ROOT + EMP-CHILD，EMP-OTHER 被收窄掉


def test_statistics_matches_list_for_global_role(dept_other, three_employees):
    """全局角色的 statistics 仍是全量（三态语义 None 不限）。"""
    user = _make_requester("AUD-1", EmployeeRole.AUDITOR, dept_other)
    listed = _jobcodes(_get(user, LIST_URL))
    stats = _get(user, STATISTICS_URL).data["data"]
    assert stats["total_employees"] == len(listed) == 4


# --------------------------------------------------------------------------
# export：导出行集合必须恒等于列表行集合（BF-047 不变量 × RBAC）
# --------------------------------------------------------------------------


def test_export_rows_equal_list_rows(dept_root, three_employees):
    user = _make_requester("MGR-ROOT", EmployeeRole.DEPT_MANAGER, dept_root)
    assert _export_jobcodes(user) == _jobcodes(_get(user, LIST_URL)) == {"MGR-ROOT", "EMP-ROOT", "EMP-CHILD"}


# --------------------------------------------------------------------------
# search：关键词不得成为跨部门泄露通道
# --------------------------------------------------------------------------


def test_search_keyword_cannot_reach_out_of_scope_rows(dept_root, three_employees):
    """无关部门员工的姓名即使被精确命中，越权调用者也不得检索到。"""
    user = _make_requester("MGR-ROOT", EmployeeRole.DEPT_MANAGER, dept_root)
    assert _jobcodes(_get(user, f"{SEARCH_URL}?keyword=无关员工")) == set()
    # 同一关键词在全局角色下可见，证明断言的是权限而非关键词失效
    auditor = _make_requester("AUD-2", EmployeeRole.AUDITOR, dept_root)
    assert _jobcodes(_get(auditor, f"{SEARCH_URL}?keyword=无关员工")) == {"EMP-OTHER"}


# --------------------------------------------------------------------------
# 三个旁路：原先绕过取行口径的端点必须同口径
# --------------------------------------------------------------------------


def test_active_employees_is_scoped(dept_root, three_employees):
    """/active-employees/ 原为裸 get_active_employees()（全公司批量 PII），现按范围收窄。"""
    user = _make_requester("MGR-ROOT", EmployeeRole.DEPT_MANAGER, dept_root)
    assert _jobcodes(_get(user, ACTIVE_URL)) == {"MGR-ROOT", "EMP-ROOT", "EMP-CHILD"}


def test_retrieve_out_of_scope_employee_returns_404(dept_root, three_employees):
    """detail 端点（lookup_field=employee_jobcode）原先仅声明式筛选，工号可枚举越权读。"""
    user = _make_requester("MGR-ROOT", EmployeeRole.DEPT_MANAGER, dept_root)
    scoped = reverse("employees-detail", kwargs={"employee_jobcode": "EMP-CHILD"})
    assert _get(user, scoped).status_code == 200
    out_of_scope = reverse("employees-detail", kwargs={"employee_jobcode": "EMP-OTHER"})
    assert _get(user, out_of_scope).status_code == 404


def test_get_employee_by_jobcode_out_of_scope_returns_404(dept_root, three_employees):
    """/employees/{jobcode}/ 旁路原先用 self.queryset（未收窄）。"""
    user = _make_requester("MGR-ROOT", EmployeeRole.DEPT_MANAGER, dept_root)
    url = reverse("employees-get-employee-by-jobcode", kwargs={"employee_jobcode": "EMP-OTHER"})
    assert _get(user, url).status_code == 404


def test_by_auth_user_scoped_both_ways(dept_root, three_employees):
    """/by-auth-user/{auth_id}/ 旁路原先直连 Employee.objects，ID 可枚举。

    正反两面：范围内可见（并锚定 URL 捕获组 str 的 PK 强转行为），范围外 404。
    """
    in_scope = Employee.objects.get(employee_jobcode="EMP-ROOT")
    out_of_scope = Employee.objects.get(employee_jobcode="EMP-OTHER")
    auth_ids = {}
    for employee in (in_scope, out_of_scope):
        auth_user = User.objects.create_user(
            auth_username=f"AUTH-{employee.employee_jobcode}",
            password=TEST_PASSWORD,
            auth_phone=_phone(),
        )
        Employee.objects.filter(pk=employee.pk).update(auth_user=auth_user)
        auth_ids[employee.employee_jobcode] = auth_user.pk

    caller = _make_requester("MGR-ROOT", EmployeeRole.DEPT_MANAGER, dept_root)

    ok_url = reverse("employees-by-auth-user", kwargs={"auth_id": auth_ids["EMP-ROOT"]})
    response = _get(caller, ok_url)
    assert response.status_code == 200, response.status_code
    assert response.data["data"]["employee_jobcode"] == "EMP-ROOT"

    denied_url = reverse("employees-by-auth-user", kwargs={"auth_id": auth_ids["EMP-OTHER"]})
    assert _get(caller, denied_url).status_code == 404


def test_by_auth_user_with_malformed_id_returns_404(dept_root, three_employees):
    """畸形 auth_id 收敛为 404：原实现会 ValueError -> 500。"""
    caller = _make_requester("MGR-ROOT", EmployeeRole.DEPT_MANAGER, dept_root)
    url = reverse("employees-by-auth-user", kwargs={"auth_id": "not-a-number"})
    assert _get(caller, url).status_code == 404
