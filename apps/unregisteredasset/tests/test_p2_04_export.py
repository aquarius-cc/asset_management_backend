"""P2-04: UnregisteredAsset 挂 ExportExcelMixin 后 /export/ 端点回归。

xlsx 行集断言先例: test_operation_log_export_scope.py。
"""

import io

import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.unregisteredasset.models import UnregisteredAsset
from apps.unregisteredasset.views import UnregisteredAssetViewSet
from apps.usermanagement.models import Employee, EmployeeRole
from core.tests import TEST_PASSWORD


User = get_user_model()

pytestmark = pytest.mark.django_db(transaction=True)


def _phone_seq():
    n = 0
    while True:
        n += 1
        yield f"140{n:08d}"


_phones = _phone_seq()


def _make_user(username, role=EmployeeRole.ASSET_ADMIN):
    user = User.objects.create_user(
        auth_username=username,
        password=TEST_PASSWORD,
        auth_is_staff=(role == EmployeeRole.SYSTEM_ADMIN),
        auth_phone=next(_phones),
    )
    Employee.objects.create(
        employee_jobcode=username,
        employee_name=f"{username}员工",
        role=role,
        employee_phone=next(_phones),
    )
    return user


def _make_unregistered(code, employee, name="测试未登记资产"):
    return UnregisteredAsset.objects.create(
        unregistered_code=code,
        scenario_type="s1_no_record",
        discovery_date="2024-07-01",
        discovery_location="测试地点",
        discovery_person=employee,
        asset_name=name,
    )


def _export_first_column(user, sheet_name="未登记资产"):
    url = reverse("unregisteredasset:unregisteredasset-export-excel")
    request = APIRequestFactory().get(url)
    force_authenticate(request, user=user)
    response = UnregisteredAssetViewSet.as_view({"get": "export_excel"})(request)
    assert response.status_code == 200, getattr(response, "data", response.status_code)
    try:
        payload = b"".join(response.streaming_content)
    finally:
        response.close()
    sheet = load_workbook(io.BytesIO(payload))[sheet_name]
    return {row[0] for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True) if row[0] is not None}


class TestUnregisteredAssetExport:
    def test_export_returns_xlsx_rows(self, db):
        employee = Employee.objects.create(
            employee_jobcode="UNR-EXP-E1",
            employee_name="发现人1",
            role=EmployeeRole.REGULAR_USER,
            employee_phone=next(_phones),
        )
        _make_unregistered("UNR-EXP-001", employee)
        user = _make_user("exp-unr-admin", EmployeeRole.SYSTEM_ADMIN)
        codes = _export_first_column(user)
        assert "UNR-EXP-001" in codes

    def test_export_regular_user_403(self, db):
        employee = Employee.objects.create(
            employee_jobcode="UNR-EXP-E2",
            employee_name="发现人2",
            role=EmployeeRole.REGULAR_USER,
            employee_phone=next(_phones),
        )
        _make_unregistered("UNR-EXP-002", employee)
        user = _make_user("exp-unr-regular", EmployeeRole.REGULAR_USER)
        url = reverse("unregisteredasset:unregisteredasset-export-excel")
        request = APIRequestFactory().get(url)
        force_authenticate(request, user=user)
        response = UnregisteredAssetViewSet.as_view({"get": "export_excel"})(request)
        assert response.status_code == 403
