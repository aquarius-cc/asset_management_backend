"""core.excel_export 内核测试（真流式导出 + XML 列宽注入）。

覆盖：
1. widths: CJK 宽度计算、``<cols>`` 生成、工作表 part 动态解析、注入后可被 openpyxl 回读
2. workbook: write_only 落盘、display_map 映射、嵌套字段、返回列宽
3. streaming: EXPORT_MAX_ROWS 守卫、limit/offset 语义与校验、响应头、临时目录清理
"""

import io
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from django.test import override_settings
from openpyxl import Workbook, load_workbook

from core.excel_export import (
    COLUMN_PADDING,
    MAX_COLUMN_WIDTH,
    ExportPaginationError,
    ExportTooLargeError,
    MissingColumnsError,
    apply_column_widths,
    build_cols_xml,
    build_excel_export_response,
    calculate_display_width,
    parse_export_bounds,
    write_export_workbook,
)
from core.excel_export.widths import resolve_worksheet_part


COLUMNS = [
    {"header": "资产编码", "field": "asset_code"},
    {"header": "状态", "field": "status", "display_map": {"in_store": "在库"}},
]
BOUND_MAX = 100


class _FakeQueryset:
    """提供 count/iterator 切片语义的最小仿真 queryset"""

    def __init__(self, rows):
        self._rows = rows

    def count(self):
        return len(self._rows)

    def iterator(self, chunk_size=None):
        return iter(self._rows)

    def __getitem__(self, key):
        return _FakeQueryset(self._rows[key])

    def __iter__(self):
        return iter(self._rows)


def _rows(n):
    return [SimpleNamespace(asset_code=f"A{i:03d}", status="in_store") for i in range(n)]


def _read_xlsx(response):
    """消费流式响应体并返回 xlsx 字节。"""
    return b"".join(response.streaming_content)


# --------------------------------------------------------------------------
# widths
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "value,expected",
    [
        (None, 0),
        ("", 0),
        ("abc", 3),
        ("资产", 4),
        ("资产A", 5),
        # U+3000 表意空格为 Fullwidth：2 + 在(2) + 库(2) + 2 = 8
        ("　在库　", 8),
    ],
)
def test_calculate_display_width(value, expected):
    assert calculate_display_width(value) == expected


def test_build_cols_xml_empty_returns_empty_string():
    assert build_cols_xml({}) == ""


def test_build_cols_xml_sorted_padded_and_capped():
    xml = build_cols_xml({2: 5, 1: 100})
    assert xml.startswith("<cols>") and xml.endswith("</cols>")
    assert xml.index('min="1"') < xml.index('min="2"')
    assert f'width="{MAX_COLUMN_WIDTH}"' in xml
    assert f'width="{5 + COLUMN_PADDING}"' in xml
    assert 'customWidth="1"' in xml


def test_resolve_worksheet_part_dynamic_not_hardcoded():
    buffer = io.BytesIO()
    workbook = Workbook()
    workbook.active.title = "首表"
    workbook.create_sheet("次表")
    workbook.save(buffer)
    with zipfile.ZipFile(buffer) as archive:
        part = resolve_worksheet_part(archive)
        assert part in archive.namelist()
    assert part == "xl/worksheets/sheet1.xml"


def test_resolve_worksheet_part_missing_file_raises():
    with zipfile.ZipFile(io.BytesIO(), "w") as archive:
        with pytest.raises(ValueError, match=r"workbook\.xml"):
            resolve_worksheet_part(archive)


def test_apply_column_widths_injected_and_readable(tmp_path):
    source, target = tmp_path / "raw.xlsx", tmp_path / "out.xlsx"
    write_export_workbook(target_path=source, rows=_rows(3), columns=COLUMNS, sheet_name="数据导出")
    with zipfile.ZipFile(source) as archive:
        assert "<cols" not in archive.read("xl/worksheets/sheet1.xml").decode("utf-8")

    apply_column_widths(source, target, {1: 6, 2: 4})
    with zipfile.ZipFile(target) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "<cols>" in sheet_xml
    assert sheet_xml.index("<cols>") < sheet_xml.index("<sheetData>")

    sheet = load_workbook(target)["数据导出"]
    assert sheet.column_dimensions["A"].width == 6 + COLUMN_PADDING
    assert [cell.value for cell in sheet[1]] == ["资产编码", "状态"]


def test_apply_column_widths_no_widths_copies_bytes(tmp_path):
    source, target = tmp_path / "raw.xlsx", tmp_path / "out.xlsx"
    write_export_workbook(target_path=source, rows=_rows(1), columns=COLUMNS, sheet_name="s")
    apply_column_widths(source, target, {})
    assert target.read_bytes() == source.read_bytes()


# --------------------------------------------------------------------------
# workbook
# --------------------------------------------------------------------------


def test_write_export_workbook_applies_display_map_and_nested_field(tmp_path):
    columns = [
        {"header": "编码", "field": "rel__code"},
        {"header": "状态", "field": "status", "display_map": {"in_store": "在库"}},
    ]
    rows = [SimpleNamespace(rel=SimpleNamespace(code="X1"), status="in_store")]
    widths = write_export_workbook(target_path=tmp_path / "o.xlsx", rows=rows, columns=columns, sheet_name="s")
    sheet = load_workbook(tmp_path / "o.xlsx")["s"]
    assert [cell.value for cell in sheet[2]] == ["X1", "在库"]
    # 表头"编码"=4 宽，值"X1"=2 宽 -> 取 4；"状态"/"在库" 均 4 宽
    assert widths == {1: 4, 2: 4}


def test_write_export_workbook_header_counts_toward_width(tmp_path):
    widths = write_export_workbook(target_path=tmp_path / "o.xlsx", rows=[], columns=COLUMNS, sheet_name="s")
    assert widths == {1: 8, 2: 4}


# --------------------------------------------------------------------------
# streaming
# --------------------------------------------------------------------------


def test_build_response_within_limit(tmp_path):
    with override_settings(EXPORT_MAX_ROWS=10):
        response = build_excel_export_response(
            rows=_FakeQueryset(_rows(3)),
            columns=COLUMNS,
            filename="a.xlsx",
            sheet_name="s",
        )
        assert response.status_code == 200
        assert response["Content-Type"].startswith("application/vnd.openxmlformats")
        assert response["X-Export-Max-Rows"] == "10"
        assert response["X-Export-Total-Count"] == "3"
        payload = _read_xlsx(response)
    assert zipfile.is_zipfile(io.BytesIO(payload))


def test_build_response_over_limit_raises():
    with override_settings(EXPORT_MAX_ROWS=2), pytest.raises(ExportTooLargeError) as exc:
        build_excel_export_response(rows=_FakeQueryset(_rows(3)), columns=COLUMNS, filename="a.xlsx", sheet_name="s")
    assert exc.value.total == 3
    assert exc.value.max_rows == 2
    assert "超过上限2" in exc.value.message


def test_build_response_missing_columns_raises():
    with pytest.raises(MissingColumnsError):
        build_excel_export_response(rows=_FakeQueryset([]), columns=[], filename="a.xlsx", sheet_name="s")


def test_build_response_throttled_skips_total_rejection():
    with override_settings(EXPORT_MAX_ROWS=2):
        response = build_excel_export_response(
            rows=_FakeQueryset(_rows(10)),
            columns=COLUMNS,
            filename="a.xlsx",
            sheet_name="s",
            limit=2,
        )
        assert response["X-Export-Total-Count"] == "2"
        _read_xlsx(response)


@pytest.mark.django_db
def test_streaming_response_close_removes_temp_dir():
    """close() 需关句柄后再删目录。

    需 django_db：``HttpResponseBase.close()`` 会发送 ``request_finished``
    信号触发 DB 连接清理，在无标记的单测里会被 pytest-django 拦截。
    """
    with override_settings(EXPORT_MAX_ROWS=10):
        response = build_excel_export_response(
            rows=_FakeQueryset(_rows(1)), columns=COLUMNS, filename="a.xlsx", sheet_name="s"
        )
        temp_dir = response._temp_dir
        assert Path(temp_dir).exists()
        _read_xlsx(response)
        response.close()
    assert not Path(temp_dir).exists()


# --------------------------------------------------------------------------
# limit/offset 语义
# --------------------------------------------------------------------------


@pytest.mark.parametrize(
    "params,expected",
    [
        ({}, (None, None)),
        ({"limit": "5"}, (5, None)),
        ({"limit": "5", "offset": "10"}, (5, 10)),
        ({"offset": "0"}, (None, 0)),
        ({"limit": "999999"}, (BOUND_MAX, None)),
    ],
)
def test_parse_export_bounds(params, expected):
    assert parse_export_bounds(params, BOUND_MAX) == expected


@pytest.mark.parametrize(
    "params,message",
    [
        ({"limit": "abc"}, "limit 必须为整数"),
        ({"limit": "0"}, "limit 必须大于等于 1"),
        ({"offset": "-1"}, "offset 必须大于等于 0"),
        ({"offset": "x"}, "offset 必须为整数"),
    ],
)
def test_parse_export_bounds_invalid(params, message):
    with pytest.raises(ExportPaginationError, match=message):
        parse_export_bounds(params, 100)
