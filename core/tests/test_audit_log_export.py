"""P2-04: AuditLogExportView 回归。

覆盖:
- 端到端 xlsx 导出 + 与列表共用 parse_audit_log_filters(筛选生效)
- RBAC: regular_user 403, auditor 200
- 参数非法时 400(与列表同文案)

先例: test_audit_log_api.py(用户/日志夹具) + test_operation_log_export_scope.py(xlsx 断言)。
"""

import io

import pytest
from django.contrib.auth import get_user_model
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.usermanagement.models import Employee, EmployeeRole
from core.audit_log_views import AuditLogExportView
from core.models_audit import AuditLog
from core.tests import TEST_PASSWORD


User = get_user_model()

pytestmark = pytest.mark.django_db(transaction=True)


def _phone_seq():
    n = 0
    while True:
        n += 1
        yield f"141{n:08d}"


_phones = _phone_seq()


def _make_user(username, role):
    user = User.objects.create_user(
        auth_username=username,
        password=TEST_PASSWORD,
        auth_is_staff=(role == EmployeeRole.SYSTEM_ADMIN),
        auth_phone=next(_phones),
    )
    Employee.objects.create(
        employee_jobcode=username,
        employee_name=f"{username}员工",
        role=role,
        employee_phone=next(_phones),
    )
    return user


def _seed_logs():
    AuditLog.objects.create(
        record_code="DEPT-EXP-1",
        app_label="department",
        operation_type="create",
        description="创建部门:导出测试",
        operator_jobcode="AUD-EXP",
        operator_name="审计员",
        ip_address="10.0.0.1",
    )
    AuditLog.objects.create(
        record_code="EMP-EXP-1",
        app_label="employee",
        operation_type="update",
        description="更新员工:导出测试",
        operator_jobcode="AUD-EXP2",
        operator_name="操作人2",
        ip_address="10.0.0.2",
    )


def _export_record_codes(user, query=""):
    url = reverse("audit-log-export") + query
    request = APIRequestFactory().get(url)
    force_authenticate(request, user=user)
    response = AuditLogExportView.as_view()(request)
    assert response.status_code == 200, getattr(response, "data", response.status_code)
    try:
        payload = b"".join(response.streaming_content)
    finally:
        response.close()
    sheet = load_workbook(io.BytesIO(payload))["审计日志"]
    return {row[0] for row in sheet.iter_rows(min_row=2, max_col=1, values_only=True) if row[0] is not None}


class TestAuditLogExport:
    def test_export_all_logs(self, db):
        _seed_logs()
        user = _make_user("exp-aud-auditor", EmployeeRole.AUDITOR)
        codes = _export_record_codes(user)
        assert codes == {"DEPT-EXP-1", "EMP-EXP-1"}

    def test_export_respects_app_label_filter(self, db):
        """与列表共用 parse_audit_log_filters: app_label 筛选生效。"""
        _seed_logs()
        user = _make_user("exp-aud-filter", EmployeeRole.AUDITOR)
        codes = _export_record_codes(user, "?app_label=department")
        assert codes == {"DEPT-EXP-1"}
        assert "EMP-EXP-1" not in codes

    def test_export_invalid_operation_type_400(self, db):
        user = _make_user("exp-aud-bad", EmployeeRole.AUDITOR)
        url = reverse("audit-log-export") + "?operation_type=not_a_real_type"
        request = APIRequestFactory().get(url)
        force_authenticate(request, user=user)
        response = AuditLogExportView.as_view()(request)
        assert response.status_code == 400
        assert "无效的操作类型" in response.data["message"]

    def test_export_regular_user_403(self, db):
        _seed_logs()
        user = _make_user("exp-aud-regular", EmployeeRole.REGULAR_USER)
        url = reverse("audit-log-export")
        request = APIRequestFactory().get(url)
        force_authenticate(request, user=user)
        response = AuditLogExportView.as_view()(request)
        assert response.status_code == 403

    def test_export_system_admin_200(self, db):
        _seed_logs()
        user = _make_user("exp-aud-sysadmin", EmployeeRole.SYSTEM_ADMIN)
        codes = _export_record_codes(user)
        assert len(codes) == 2
